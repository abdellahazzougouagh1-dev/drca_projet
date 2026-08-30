import os
from docx import Document

def inspect_docx(path, out_path):
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(f"--- {path} ---\n")
        doc = Document(path)
        for i, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if text:
                f.write(f"Para {i}: {text}\n")
        f.write("\nTables:\n")
        for i, table in enumerate(doc.tables):
            f.write(f"Table {i}:\n")
            for row in table.rows:
                f.write(" | ".join([cell.text.strip() for cell in row.cells]) + "\n")

inspect_docx(r'storage\app\templates\resultats_ao.docx', 'rao_inspect.txt')
